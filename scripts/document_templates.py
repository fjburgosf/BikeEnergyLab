"""Author on the retained template components, preserving its visual system."""

import json
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / "docs/templates/catalog.json"
AUTHORS = "Francisco Javier Burgos Flórez\nJuan Guillermo Popayán Hernández"
REVISION = "Revisión de entrega del 8 de octubre de 2026"


def replace_text(paragraph, text):
    rpr = next((deepcopy(r._r.rPr) for r in paragraph.runs if r._r.rPr is not None), None)
    paragraph.clear()
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def append_role(doc, prototype, text):
    element = deepcopy(prototype)
    doc.element.body.insert(len(doc.element.body) - 1, element)
    p = Paragraph(element, doc._body)
    replace_text(p, text)
    return p


def section_break(doc, index):
    p = doc.add_paragraph()
    p._p.get_or_add_pPr().append(deepcopy(doc._template_sections[index]))


def data_table(doc, rows, widths=None, *, metadata=False):
    candidates = doc._template_tables
    prototype = (
        candidates[0]
        if metadata
        else next(
            (
                t
                for t in candidates[1:]
                if len(Table(t, doc._body).columns) == len(rows[0])
                and len(Table(t, doc._body).rows) > 1
            ),
            candidates[0],
        )
    )
    element = deepcopy(prototype)
    tr = element.findall(qn("w:tr"))
    patterns = [deepcopy(tr[0]), deepcopy(tr[min(1, len(tr) - 1)])]
    for row in tr:
        element.remove(row)
    if len(Table(prototype, doc._body).columns) != len(rows[0]):
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table._tbl.tblPr.clear()
        for prop in list(prototype.find(qn("w:tblPr"))):
            table._tbl.tblPr.append(deepcopy(prop))
        for _ in rows:
            table.add_row()
        element = table._tbl
    else:
        doc.element.body.insert(len(doc.element.body) - 1, element)
        for i in range(len(rows)):
            element.append(deepcopy(patterns[0 if i == 0 and not metadata else 1]))
        table = Table(element, doc._body)
    table.autofit = False
    if widths is None and not metadata:
        widths = [2.15, 4.35] if len(rows[0]) == 2 else [6.5 / len(rows[0])] * len(rows[0])
    if widths is not None:
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
    for i, (row, values) in enumerate(zip(table.rows, rows)):
        props = row._tr.get_or_add_trPr()
        if props.find(qn("w:cantSplit")) is None:
            props.append(OxmlElement("w:cantSplit"))
        if i == 0 and not metadata and props.find(qn("w:tblHeader")) is None:
            props.append(OxmlElement("w:tblHeader"))
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            if widths is not None:
                cell.width = Inches(widths[j])
            paragraph = cell.paragraphs[0]
            for extra in list(cell.paragraphs)[1:]:
                extra._p.getparent().remove(extra._p)
            replace_text(paragraph, str(value))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_with_next = False
            for run in paragraph.runs:
                run.bold = (j == 0) if metadata else i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def template_document(title, short, introduction, version):
    kind = (
        "user"
        if short == "Manual de usuario"
        else "description"
        if short == "Descripción del software"
        else "technical"
    )
    catalog = json.loads(CATALOG.read_text(encoding="utf-8"))
    entry = catalog[kind]
    doc = Document(ROOT / entry["master"])
    doc._template_sections = [deepcopy(s._sectPr) for s in doc.sections]
    doc._template_tables = [deepcopy(t._tbl) for t in doc.tables]
    roles = {
        key: deepcopy(doc.paragraphs[entry[key]]._p)
        for key in ["title_paragraph", "subtitle_paragraph", "introduction_paragraph"]
    }
    divider = (
        deepcopy(doc.paragraphs[entry["divider_paragraph"]]._p)
        if entry["divider_paragraph"] is not None
        else None
    )
    for element in list(doc.element.body):
        if element.tag != qn("w:sectPr"):
            doc.element.body.remove(element)
    doc._template_kind, doc._chapter_count, doc._figure_count = kind, 0, 0
    for part in doc.part.related_parts.values():
        if str(part.partname).startswith("/word/header"):
            for element in part.element.xpath(".//w:p"):
                p = Paragraph(element, doc._body)
                if p._p.xpath(".//w:pBdr") or p.runs:
                    replace_text(p, f"BikeEnergyLab — {short}")
    append_role(doc, roles["title_paragraph"], "BikeEnergyLab")
    append_role(doc, roles["subtitle_paragraph"], short)
    if divider is not None:
        doc.element.body.insert(len(doc.element.body) - 1, divider)
    append_role(doc, roles["introduction_paragraph"], introduction)
    metadata = [
        ["Autores", AUTHORS],
        ["Versión", version],
        ["Fecha", "7 de octubre de 2026"],
        ["Revisión", REVISION],
        ["Contacto", "fjburgosf@gmail.com"],
    ]
    data_table(doc, metadata, metadata=True)
    if kind == "description":
        doc.add_page_break()
    else:
        section_break(doc, 0)
    p = doc.add_paragraph("Tabla de contenido")
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(16)
    if kind == "description":
        p.runs[0].font.color.rgb = doc.styles["Heading 1"].font.color.rgb
    p = doc.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\h \\o "1-2" \\z \\u')
    p._p.append(field)
    if kind == "description":
        doc.add_page_break()
    else:
        section_break(doc, 1)
    if "Code" not in doc.styles:
        code = doc.styles.add_style("Code", 1)
        code.base_style = doc.styles["Normal"]
        code.font.name, code.font.size = "Consolas", Pt(9)
    doc.core_properties.title = title
    doc.core_properties.author = AUTHORS.replace("\n", " y ")
    doc.core_properties.subject = f"BikeEnergyLab {version} — {REVISION}"
    doc.core_properties.language = "es-CO"
    return doc


def chapter(doc, heading, level):
    if level == 1:
        if doc._chapter_count and doc._template_kind != "description":
            index = min(doc._chapter_count + 1, len(doc._template_sections) - 2)
            section_break(doc, index)
        doc._chapter_count += 1
    p = doc.add_heading(heading, level)
    if level == 1 and doc._template_kind == "description" and doc._chapter_count > 1:
        p.paragraph_format.page_break_before = True
    return p


def figure(doc, relative, caption):
    path = ROOT / relative
    if not path.is_file():
        raise FileNotFoundError(f"Capture required: {path}")
    doc._figure_count += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(6.5))
    picture._inline.docPr.set("descr", caption)
    p = doc.add_paragraph()
    p.add_run(f"Figura {doc._figure_count}. ").bold = True
    p.add_run(caption)


def functions_document(content, version):
    entry = json.loads(CATALOG.read_text(encoding="utf-8"))["functions"]
    doc = Document(ROOT / entry["master"])
    replace_text(doc.paragraphs[0], f"BikeEnergyLab {version}")
    replace_text(doc.paragraphs[1], content)
    doc.core_properties.title = "Título y descripción de funciones de BikeEnergyLab"
    doc.core_properties.author = AUTHORS.replace("\n", " y ")
    doc.core_properties.subject = REVISION
    return doc
