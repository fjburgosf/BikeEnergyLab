"""Build five template-based Spanish documents with editable Word math.

Run with the Python returned by load_workspace_dependencies, then render every page.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import tomllib
from datetime import date
from pathlib import Path

from document_templates import chapter, data_table, figure, functions_document, template_document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/docx"
VERSION = tomllib.loads((ROOT / "pyproject.toml").read_text(encoding="utf-8"))["project"]["version"]
RELEASE_DATE = date.fromisoformat(
    re.search(
        r"^date-released: (\d{4}-\d{2}-\d{2})$",
        (ROOT / "CITATION.cff").read_text(encoding="utf-8"),
        re.MULTILINE,
    )[1]
)
BLACK = RGBColor(0, 0, 0)


def sub(base, index):
    return ("sub", base, index)


def sup(base, power):
    return ("sup", base, power)


def frac(top, bottom):
    return ("frac", top, bottom)


def math_node(token):
    if isinstance(token, str):
        run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = token
        run.append(text)
        return run
    kind, a, b = token
    if kind in {"sub", "sup"}:
        node = OxmlElement("m:sSub" if kind == "sub" else "m:sSup")
        first, second = "m:e", "m:sub" if kind == "sub" else "m:sup"
    elif kind == "frac":
        node = OxmlElement("m:f")
        first, second = "m:num", "m:den"
    else:
        raise ValueError(kind)
    for key, items in [(first, a), (second, b)]:
        child = OxmlElement(key)
        for item in items if isinstance(items, list) else [items]:
            child.append(math_node(item))
        node.append(child)
    return node


def add_equation(doc, tokens):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    math = OxmlElement("m:oMath")
    for token in tokens:
        math.append(math_node(token))
    p._p.append(math)


EQUATIONS = {
    "forces": [
        [
            "θ = atan(grade)     m = ",
            sub("m", "bike"),
            " + ",
            sub("m", "rider"),
            " + ",
            sub("m", "cargo"),
        ],
        [
            sub("F", "roll"),
            " = mg",
            sub("C", "rr"),
            "cos(θ)     ",
            sub("F", "grade"),
            " = mg sin(θ)",
        ],
        [
            sub("v", "rel"),
            " = v − w     ",
            sub("F", "aero"),
            " = ",
            frac("1", "2"),
            "ρ",
            sub("C", "d"),
            "A",
            sub("v", "rel"),
            "|",
            sub("v", "rel"),
            "|",
        ],
        [sub("F", "acc"), " = ma"],
    ],
    "wheel": [
        [
            sub("P", "wheel"),
            " = (",
            sub("F", "roll"),
            " + ",
            sub("F", "grade"),
            " + ",
            sub("F", "aero"),
            " + ",
            sub("F", "acc"),
            ")v",
        ]
    ],
    "density": [["ρ = ", frac("p", "287.05(T + 273.15)"), "     T en °C"]],
    "power": [
        [
            sub("P", "human wheel"),
            " = ",
            sub("η", "drive"),
            sub("P", "pedal"),
            "     limitado a demanda positiva",
        ],
        [
            sub("P", "shaft"),
            " = ",
            frac(sub("P", "motor wheel"), sub("η", "drive")),
            "     ",
            sub("P", "electric motor"),
            " = ",
            frac(sub("P", "shaft"), sub("η", "motor")),
        ],
    ],
    "capacity": [
        [sub("E", "full"), " = ", sub("E", "nominal"), sub("f", "usable"), sub("f", "temperature")],
        [
            sub("E", "available"),
            " = ",
            sub("E", "full"),
            "(",
            sub("SOC", "initial"),
            " − ",
            sub("SOC", "min"),
            ")",
        ],
    ],
    "soc": [
        [
            "ΔSOC = −",
            frac([sub("P", "terminal"), "Δt"], ["3600", sub("E", "full")]),
            "     Δt en s y E en Wh",
        ]
    ],
    "ecm": [
        ["V = OCV(SOC) − ", sub("R", "0"), "I − ", sub("V", "RC")],
        [
            frac(["d", sub("V", "RC")], "dt"),
            " = −",
            frac(sub("V", "RC"), [sub("R", "1"), sub("C", "1")]),
            " + ",
            frac("I", sub("C", "1")),
        ],
        ["ΔSOC = −", frac("IΔt", ["3600", sub("Q", "effective")]), "     Q en Ah"],
    ],
    "terminal": [
        [
            sub("P", "terminal"),
            " = ",
            sub("P", "electric motor"),
            " + ",
            sub("P", "aux"),
            " − ",
            sub("P", "regen"),
        ]
    ],
    "fit": [
        [
            "θ* = ",
            sub("arg min", "θ"),
            sub("Σ", "i"),
            sub("w", "i"),
            sup(["(", sub("E", "pred i"), "(θ) − ", sub("E", "obs i"), ")"], "2"),
        ]
    ],
    "residual": [
        [
            sub("r", "i"),
            " = ",
            frac([sub("E", "obs i"), " − ", sub("E", "phys i")], sub("L", "i")),
            "     r en Wh/km",
        ]
    ],
    "models": [
        ["M1     Ê = ", sub("E", "phys"), "        M2     Ê = ", sub("E", "data")],
        ["M3     Ê = ", sub("E", "phys"), " + r̂L"],
        ["M4     Ê = ", sub("E", "phys"), " + α(x)r̂L"],
    ],
    "gate": [
        [
            "α(x) = ",
            frac("n", "n + 20"),
            frac(["exp(−", sup("max(d − 1, 0)", "2"), ")"], ["1 + ", sup("u", "2")]),
        ]
    ],
    "conformal": [
        ["k = ⌈(n + 1)c⌉"],
        [sub("I", "c"), " = [Ê − ", sub("q", "c"), "L, Ê + ", sub("q", "c"), "L]"],
    ],
    "mission": [["p̂ = ", frac([sub("N", "successful")], "N")]],
    "sobol": [
        [
            sub("S", "i"),
            " = ",
            frac(["mean[(f(B) − f̄)(f(", sub("AB", "i"), ") − f(A))]"], "Var(f)"),
        ],
        [
            sub("ST", "i"),
            " = ",
            frac(["mean[", sup(["(f(A) − f(", sub("AB", "i"), "))"], "2"), "]"], "2 Var(f)"),
        ],
    ],
}


def field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in [begin, code, end]:
        run = OxmlElement("w:r")
        run.append(element)
        paragraph._p.append(run)


def inline(paragraph, text):
    pattern = r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
    for piece in re.split(pattern, text):
        if piece.startswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif piece.startswith("[") and "](" in piece:
            label, target = piece[1:].split("](", 1)
            target = target[:-1]
            if target.startswith("http"):
                relation = paragraph.part.relate_to(
                    target,
                    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                    is_external=True,
                )
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), relation)
                run = OxmlElement("w:r")
                props = OxmlElement("w:rPr")
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "17517A")
                props.append(color)
                run.append(props)
                word = OxmlElement("w:t")
                word.text = label
                run.append(word)
                link.append(run)
                paragraph._p.append(link)
            else:
                paragraph.add_run(f"{label} ({target})")
        else:
            paragraph.add_run(piece)


def table(doc, rows, widths=None):
    return data_table(doc, rows, widths)


def documentation_prose(text):
    """Split explanatory clauses without changing packaged GUI text or labels."""
    parts = text.split("; ")
    return ". ".join(part[0].upper() + part[1:] if part else part for part in parts)


def special(doc, text):
    if text == "{{TUTORIAL}}":
        data = json.loads(
            (ROOT / "docs/images/gui-1.0.0/captures.json").read_text(encoding="utf-8")
        )
        table(
            doc,
            [["Paso y botón", "Uso"]]
            + [
                [f"{i}. {s['title']}\n{s['button']}", documentation_prose(s["description"])]
                for i, s in enumerate(data["tutorial"], 1)
            ],
            [2.2, 4.3],
        )
    elif text == "{{EXAMPLES}}":
        data = json.loads(
            (ROOT / "docs/images/gui-1.0.0/captures.json").read_text(encoding="utf-8")
        )
        for i, example in enumerate(data["examples"], 1):
            key = example["key"]
            doc.add_heading(f"4.14.{i} {example['title'].split('·')[-1].strip()}", 3)
            doc.add_paragraph(documentation_prose(example["description"]))
            doc.add_paragraph(
                "Procedimiento: seleccione este ejemplo en Inicio, pulse Cargar ejemplo, confirme el caso activo y pulse Simular. Abra Resultados al terminar."
            )
            result = data["example_results"][key]
            if i <= 5:
                values = f"Energía {result['energy_wh']:.3f} Wh, consumo {result['wh_per_km']:.3f} Wh/km y SOC final {result['final_soc']:.5f}. Ruta completada y factible en esta práctica."
            elif key == "calibration":
                values = "Se recuperaron Crr = 0.008 y CdA = 0.48 m². El rango del Jacobiano fue 2. El caso carece de ruido/discrepancia y no demuestra recuperación con datos reales."
            elif key == "hybrid":
                values = "El modelo se entrenó con 60 rutas y calibró intervalos con 25 independientes. Queda disponible en memoria para volver a simular y guardar. Revise predicción M4, alpha y soporte OOD. El SOC sigue la simulación física."
            else:
                values = f"Con {result['n_samples']} muestras, la probabilidad física de misión fue {100 * result['mission_probability']:.2f}% y la demanda media de ruta completa fue {result['full_route_demand_wh']['mean']:.3f} Wh. El intervalo Monte Carlo refleja muestreo bajo los supuestos configurados."
                if key == "uncertainty":
                    values += " La mediana de autonomía alcanza el horizonte de 80 km. Examine la fracción censurada."
                else:
                    values += " SOC inicial 0.38 y reserva 0.15. Se exige completar los 26 km con potencia y reserva. El consumo hasta agotamiento y la demanda completa pueden diferir."
            doc.add_paragraph("Resultado observado: " + values)
            figure(
                doc,
                f"docs/images/gui-1.0.0/ejemplo-{i:02d}-{key}.png",
                f"Ejecución del ejemplo {i:02d} {example['title'].split('·')[-1].strip()}.",
            )
    elif text.startswith("{{IMAGE:"):
        relative, caption = text[8:-2].split("|", 1)
        figure(doc, relative, caption)
    elif text.startswith("{{EQ:"):
        for tokens in EQUATIONS[text[5:-2]]:
            add_equation(doc, tokens)
    elif text == "{{SOURCE}}":
        source = hashlib.sha256()
        package = ROOT / "src/bikeenergylab"
        for path in sorted(package.rglob("*.py")):
            source.update(path.relative_to(package).as_posix().encode())
            source.update(path.read_bytes())
        identity = source.hexdigest()
        doc.add_paragraph(f"Identidad SHA256 de las fuentes de la versión {VERSION}")
        p = doc.add_paragraph(identity)
        p.paragraph_format.space_after = Pt(10)
        p.runs[0].font.name = "Consolas"
        p.runs[0].font.size = Pt(9.5)
    elif text == "{{TABLE:ID}}":
        path = ROOT / "docs/validation_data/predictive_aggregate.csv"
        rows = [["Modelo", "MAE Wh", "CRPS Wh", "Cobertura 95%", "Brier"]]
        with path.open(encoding="utf-8") as stream:
            for row in csv.DictReader(stream):
                if row["scenario"] == "ID":
                    rows.append(
                        [
                            row["model"],
                            f"{float(row['mae_wh_mean']):.3f}",
                            f"{float(row['crps_wh_mean']):.3f}",
                            f"{100 * float(row['empirical_coverage95_mean']):.2f}%",
                            f"{float(row['energy_budget_brier_mean']):.4f}",
                        ]
                    )
        table(doc, rows, [1.0, 1.2, 1.2, 1.6, 1.5])
    elif text == "{{TABLE:DRIFT}}":
        path = ROOT / "docs/validation_data/prequential_metrics.csv"
        rows = [["Estrategia", "Modelo", "MAE Wh", "RMSE Wh"]]
        with path.open(encoding="utf-8") as stream:
            for row in csv.DictReader(stream):
                rows.append(
                    [
                        "Adaptativa" if row["strategy"] == "adaptive" else "Congelada",
                        row["model"],
                        f"{float(row['mae']):.3f}",
                        f"{float(row['rmse']):.3f}",
                    ]
                )
        table(doc, rows, [2.0, 1.2, 1.65, 1.65])
    else:
        raise ValueError(text)


def markdown(doc, content, user_manual=False):
    lines = content.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("{{"):
            special(doc, line)
        elif line.startswith("##"):
            level = len(line) - len(line.lstrip("#")) - 1
            chapter(doc, line.lstrip("# ").strip(), min(level, 3))
        elif line.startswith("```"):
            while i < len(lines) and not lines[i].startswith("```"):
                p = doc.add_paragraph(lines[i], "Code")
                p.paragraph_format.keep_with_next = i + 1 < len(lines) and not lines[
                    i + 1
                ].startswith("```")
                i += 1
            i += 1
        elif line.startswith("|"):
            rows = [[x.strip() for x in line.strip("|").split("|")]]
            while i < len(lines) and lines[i].strip().startswith("|"):
                values = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", value) for value in values):
                    rows.append(values)
                i += 1
            table(doc, rows)
        else:
            parts = [line]
            while (
                i < len(lines)
                and lines[i].strip()
                and not re.match(r"^(#|\||```|\{\{|\d+\. |[-*] )", lines[i].strip())
            ):
                parts.append(lines[i].strip())
                i += 1
            text = " ".join(parts)
            match = re.match(r"\d+\. \*\*(.+?):\*\* (.*)", text)
            if user_manual and match:
                heading = match[1].replace("38%", "38 por ciento")
                doc.add_heading(heading, 2)
                text = match[2]
                openings = {
                    "ejecutar ": "Abra el programa con ",
                    "masas, ": "Configure masas, ",
                    "masa y potencia ": "Defina masa y potencia ",
                    "eficiencia constante/mapa y límites.": "Configure eficiencia constante o un mapa y sus límites.",
                    "Wh nominales, ": "Indique Wh nominales, ",
                    "crear un perfil ": "Cree un perfil ",
                    "configurar timestep y auxiliares; pulsar Simular.": "Configure el paso de tiempo y los auxiliares; pulse Simular.",
                    "cargar CSV ": "Cargue CSV ",
                    "entrenar CGPRA ": "Entrene CGPRA ",
                    "editar las distribuciones en YAML y Aplicar YAML.": "Edite las distribuciones en YAML y pulse Aplicar YAML.",
                    "configurar initial_soc": "Configure initial_soc",
                    "temperatura, viento ": "Configure temperatura, viento ",
                    "ejecutar benchmark ": "Ejecute benchmark ",
                    "escoger carpeta ": "Escoja una carpeta ",
                }
                for original, revised in sorted(
                    openings.items(), key=lambda item: len(item[0]), reverse=True
                ):
                    if text.startswith(original):
                        text = revised + text[len(original) :]
                        break
                text = text.replace("Usar constante,", "Use un valor constante,")
                text = text.replace(
                    "Use un valor constante, perfil en CSV o modo dependiente",
                    "Use un valor constante, un perfil en CSV o el modo dependiente",
                )
                text = text.replace("reemplazarlos por mediciones", "reemplácelos por mediciones")
                text = text.replace("o importar CSV/GPX", "o importe CSV/GPX")
                text = text.replace("Pulsar **Visualizar ruta**", "Pulse **Visualizar ruta**")
                text = text.replace(
                    "Configure temperatura, viento (negativo frontal), presión/densidad.",
                    "Configure temperatura, viento (negativo frontal) y presión o densidad.",
                )
                text = text.replace(
                    "Configure initial_soc=0.38, importar ruta",
                    "Configure initial_soc=0.38, importe una ruta",
                )
                text = text.replace("y seleccionar sensibilidad", "y seleccione sensibilidad")
                text = text.replace("no comparar MAE", "no compare MAE")
                text = text.replace(
                    "El cálculo corre en background.", "El cálculo se realiza en segundo plano."
                )
                text = text.replace(
                    "Revisar energía, SOC, distancia y `feasible`.",
                    "Revise energía, SOC, distancia, `completed_route` y `feasible`.",
                )
                text = text.replace("Inspeccionar reporte", "Inspeccione el reporte")
                text = text.replace("Leer resultados positivos", "Lea los resultados positivos")
            p = doc.add_paragraph()
            inline(p, text)


def new_document(title, short, introduction):
    return template_document(title, short, introduction, VERSION)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    specifications = [
        (
            f"Manual_de_usuario_BikeEnergyLab_{VERSION}.docx",
            "Manual de usuario de BikeEnergyLab",
            "Manual de usuario",
            "Este manual guía al usuario desde abrir el programa y practicar con ejemplos hasta configurar su bicicleta, simular, interpretar resultados, calibrar y exportar. El menú Ejemplos y el tutorial integrado permiten comenzar sin preparar archivos de datos.",
            ROOT / "docs/user_manual.md",
            True,
        ),
        (
            f"Manual_tecnico_BikeEnergyLab_{VERSION}.docx",
            "Manual técnico de BikeEnergyLab",
            "Manual técnico",
            "Este manual describe la arquitectura, las clases públicas, la configuración, los algoritmos y la reproducción de BikeEnergyLab. Está dirigido a desarrolladores e investigadores que necesitan utilizar la API, revisar decisiones de implementación y repetir las verificaciones de la distribución.",
            ROOT / "docs/docx_sources/manual_tecnico_es.md",
            False,
        ),
        (
            f"Metodologia_cientifica_BikeEnergyLab_{VERSION}.docx",
            "Metodología científica de BikeEnergyLab",
            "Metodología científica",
            "Este documento formula el modelo físico y el método híbrido, explica calibración e incertidumbre y presenta los resultados sintéticos conservados. Las ecuaciones son objetos editables de Word. Los hallazgos incluyen fallos OOD y comparaciones desfavorables, porque la evidencia disponible no demuestra superioridad universal de CGPRA.",
            ROOT / "docs/docx_sources/metodologia_es.md",
            False,
        ),
        (
            f"Descripcion_del_Software_BikeEnergyLab_{VERSION}.docx",
            "Descripción del software BikeEnergyLab",
            "Descripción del software",
            "Plataforma científica de escritorio para modelar el consumo energético, la autonomía y la probabilidad de misión de bicicletas eléctricas. Integra física, calibración, aprendizaje residual y análisis de incertidumbre con ejemplos sintéticos reproducibles.",
            ROOT / "docs/docx_sources/descripcion_es.md",
            False,
        ),
    ]
    for filename, title, short, intro, source, user in specifications:
        doc = new_document(title, short, intro)
        markdown(doc, source.read_text(encoding="utf-8"), user_manual=user)
        doc.add_heading("Autoría y condiciones de uso", 2)
        doc.add_paragraph(
            "Autores: Francisco Javier Burgos Flórez y Juan Guillermo Popayán Hernández. Contacto: fjburgosf@gmail.com. Conserve CITATION.cff para citar el software. Los derechos y condiciones de uso son los indicados en LICENSE. No se presume una licencia abierta."
        )
        path = OUT / filename
        doc.save(path)
        print(path)
    content = (
        (ROOT / "docs/docx_sources/titulo_funciones_es.md").read_text(encoding="utf-8").strip()
    )
    doc = functions_document(content, VERSION)
    path = OUT / f"Titulo_y_descripcion_de_funciones_BikeEnergyLab_{VERSION}.docx"
    doc.save(path)
    print(path)


if __name__ == "__main__":
    main()
